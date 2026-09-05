from pptx import Presentation
from pptx.util import Pt

path = r"C:\Users\Kavya\My Drive\Kavya Google Drive\14 Engineering\402 Cytutor\Cytutor\Cytutor_260414\cytutor-1-main\20CYS495Project Phase I.pptx"
prs = Presentation(path)

print(f"Size: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f}")
print(f"Slides: {len(prs.slides)}, Layouts: {len(prs.slide_layouts)}")
print()
for i, lay in enumerate(prs.slide_layouts):
    print(f"  Layout {i}: {lay.name}")
print()
for si, slide in enumerate(prs.slides):
    print(f"--- Slide {si+1} (layout: {slide.slide_layout.name}) ---")
    for shape in slide.shapes:
        pos = f"L={shape.left/914400:.2f} T={shape.top/914400:.2f} W={shape.width/914400:.2f} H={shape.height/914400:.2f}"
        print(f"  [{shape.name}] {pos}")
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                line = para.text.strip()
                if line:
                    sz = para.runs[0].font.size.pt if para.runs and para.runs[0].font.size else "?"
                    bold = para.runs[0].font.bold if para.runs else "?"
                    try:
                        col = str(para.runs[0].font.color.rgb) if para.runs and para.runs[0].font.color.type else "inherited"
                    except:
                        col = "inherited"
                    print(f"    TEXT: \"{line[:70]}\" sz={sz} bold={bold} col={col}")
        try:
            f = shape.fill
            if str(f.type) != "None":
                try:
                    print(f"    FILL: {f.type} #{f.fore_color.rgb}")
                except:
                    print(f"    FILL: {f.type}")
        except:
            pass
    print()
